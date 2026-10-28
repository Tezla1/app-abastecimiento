# Diseñado y Desarrollado por Jaider Esteban Rodríguez Chiribi y Sebastián González León
# Genera los reportes en Excel (resultados y contingencias) y evita que una celda se interprete como fórmula.

import csv

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ...application.ports import ExportadorReportes


def _celda_segura(v):
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    return v


class ExportadorReportesArchivo(ExportadorReportes):
    def exportar_csv(self, ruta: str, filas: list[dict]) -> None:
        if not filas:
            raise ValueError("No hay filas para exportar.")
        with open(ruta, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=list(filas[0].keys()))
            w.writeheader()
            for fila in filas:
                w.writerow({k: "" if v is None else _celda_segura(v) for k, v in fila.items()})

    def exportar_xlsx(self, ruta: str, titulo: str, meta: list[tuple], filas: list[dict]) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Comparación"
        ws["A1"] = titulo
        ws["A1"].font = Font(bold=True, size=14)
        fila = 3
        for k, v in meta:
            ws.cell(fila, 1, k).font = Font(bold=True)
            ws.cell(fila, 2, _celda_segura(v))
            fila += 1
        fila += 1
        columnas = list(filas[0].keys()) if filas else []
        for j, c in enumerate(columnas, 1):
            celda = ws.cell(fila, j, c)
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="1F3A68")
            celda.alignment = Alignment(horizontal="center", wrap_text=True)
        for i, datos in enumerate(filas, fila + 1):
            for j, c in enumerate(columnas, 1):
                v = datos[c]
                celda = ws.cell(i, j, "n/d" if v is None else _celda_segura(v))
                if isinstance(v, float):
                    celda.number_format = "0.00"
        ws.column_dimensions["A"].width = 42
        for col in "BCD":
            ws.column_dimensions[col].width = 28
        wb.save(ruta)

    def exportar_contingencias(self, ruta: str, titulo: str, cierres: list[dict], rutas: list[dict]) -> None:
        wb = Workbook()
        for i, (nombre, filas, vacio) in enumerate([("Contingencias", cierres, "Sin contingencias registradas."),
                                                    ("Rutas alternativas", rutas, "Ninguna ruta tuvo que cambiar.")]):
            ws = wb.active if i == 0 else wb.create_sheet()
            ws.title = nombre
            ws["A1"] = titulo
            ws["A1"].font = Font(bold=True, size=14)
            if not filas:
                ws["A3"] = vacio
                continue
            columnas = list(filas[0].keys())
            for j, c in enumerate(columnas, 1):
                celda = ws.cell(3, j, c)
                celda.font = Font(bold=True, color="FFFFFF")
                celda.fill = PatternFill("solid", fgColor="1F3A68")
                celda.alignment = Alignment(horizontal="center", wrap_text=True)
                ws.column_dimensions[celda.column_letter].width = max(16, min(48, len(c) + 8))
            for r, datos in enumerate(filas, 4):
                for j, c in enumerate(columnas, 1):
                    celda = ws.cell(r, j, "" if datos[c] is None else _celda_segura(datos[c]))
                    celda.alignment = Alignment(wrap_text=True, vertical="top")
                    ancho = ws.column_dimensions[celda.column_letter].width
                    ws.column_dimensions[celda.column_letter].width = max(ancho, min(60, len(str(datos[c])) + 2))
        wb.save(ruta)
