"""Lectura/escritura de talleres en Excel con openpyxl."""
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ...application.errors import ErrorDeNegocio
from ...application.ports import DatosExcel, ExcelGateway
from ...domain.entities import Conexion, Estacion, Pedido, RutaProceso, Taller, TipoPrenda

HOJAS = {
    "Taller": ["nombre", "area_m2", "operadores", "jornada_h", "origen_datos", "observaciones"],
    "Nodos": ["codigo", "nombre", "funcion", "x_m", "y_m", "tiempo_min_u", "puestos", "grupo"],
    "Conexiones": ["origen", "destino", "distancia_m", "ancho_m", "bidireccional"],
    "Prendas": ["nombre", "peso_unidad_kg", "unidades_por_lote", "factor_proceso"],
    "Rutas": ["prenda", "ruta", "fraccion_pct", "secuencia"],
    "Pedidos": ["prenda", "cantidad", "prioridad", "llegada_min", "vencimiento_min", "origen_datos"],
}
_VERDADERO = {"1", "si", "sí", "true", "verdadero", "x", "yes"}
_FALSO = {"0", "no", "false", "falso", ""}


def _estilo_encabezado(ws) -> None:
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F3A68")
        celda.alignment = Alignment(horizontal="center")
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(14, len(str(col[0].value)) + 4)


def _nuevo_libro() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    for nombre, cols in HOJAS.items():
        ws = wb.create_sheet(nombre)
        ws.append(cols)
        _estilo_encabezado(ws)
    return wb


class OpenpyxlGateway(ExcelGateway):
    def crear_plantilla(self, ruta: str) -> None:
        _nuevo_libro().save(ruta)

    def escribir_taller(self, ruta, taller, estaciones, conexiones, tipos, rutas, pedidos) -> None:
        wb = _nuevo_libro()
        nombres = {t.id: t.nombre for t in tipos}
        wb["Taller"].append([taller.nombre, taller.area_m2, taller.operadores, taller.jornada_h,
                             taller.origen_datos, taller.observaciones])
        for e in estaciones:
            wb["Nodos"].append([e.codigo, e.nombre, e.funcion, e.x, e.y, e.tiempo_proceso_min, e.capacidad, e.grupo])
        for c in conexiones:
            wb["Conexiones"].append([c.origen, c.destino, c.distancia_m, c.ancho_m, int(c.bidireccional)])
        for t in tipos:
            wb["Prendas"].append([t.nombre, t.peso_unidad_kg, t.unidades_por_lote, t.factor_proceso])
        for r in rutas:
            wb["Rutas"].append([nombres.get(r.tipo_prenda_id, ""), r.nombre, r.fraccion_pct, r.secuencia])
        for p in pedidos:
            wb["Pedidos"].append([nombres.get(p.tipo_prenda_id, ""), p.cantidad, p.prioridad,
                                  p.llegada_min, p.vencimiento_min, p.origen_datos])
        wb.save(ruta)

    def leer_taller(self, ruta: str) -> DatosExcel:
        try:
            wb = load_workbook(ruta, read_only=True, data_only=True)
        except Exception:
            raise ErrorDeNegocio("No se pudo abrir el archivo: ¿es un .xlsx válido?") from None
        datos = DatosExcel()
        try:
            for hoja, esperadas in HOJAS.items():
                if hoja not in wb.sheetnames:
                    datos.errores.append(f"Falta la hoja «{hoja}».")
                    continue
                filas = list(wb[hoja].iter_rows(values_only=True))
                encabezado = [str(c).strip().lower() if c is not None else "" for c in (filas[0] if filas else ())]
                faltan = [c for c in esperadas if c not in encabezado]
                if faltan:
                    datos.errores.append(f"Hoja «{hoja}»: faltan las columnas {', '.join(faltan)}.")
                    continue
                for n, fila in enumerate(filas[1:], 2):
                    if all(v is None or str(v).strip() == "" for v in fila):
                        continue
                    reg = {encabezado[i]: v for i, v in enumerate(fila) if i < len(encabezado)}
                    try:
                        self._cargar_fila(hoja, reg, datos)
                    except (ValueError, TypeError) as exc:
                        datos.errores.append(f"{hoja} fila {n}: {exc}")
        finally:
            wb.close()
        return datos

    # --- helpers de lectura ---
    @staticmethod
    def _txt(v) -> str:
        return "" if v is None else str(v).strip()

    @staticmethod
    def _num(v, nombre: str, defecto=None) -> float:
        if v is None or str(v).strip() == "":
            if defecto is None:
                raise ValueError(f"«{nombre}» es obligatorio.")
            return defecto
        try:
            return float(str(v).replace(",", "."))
        except ValueError:
            raise ValueError(f"«{nombre}» debe ser numérico (valor: {v!r}).") from None

    def _cargar_fila(self, hoja: str, r: dict, datos: DatosExcel) -> None:
        t, n = self._txt, self._num
        if hoja == "Taller":
            if datos.taller is not None:
                raise ValueError("solo se admite un taller por archivo.")
            datos.taller = Taller(
                nombre=t(r["nombre"]), area_m2=n(r["area_m2"], "area_m2", 0.0),
                operadores=int(n(r["operadores"], "operadores", 0)), jornada_h=n(r["jornada_h"], "jornada_h", 8.0),
                origen_datos=t(r["origen_datos"]).upper() or "SINTETICO", observaciones=t(r["observaciones"]))
        elif hoja == "Nodos":
            datos.estaciones.append(Estacion(
                0, t(r["codigo"]), t(r["nombre"]), t(r["funcion"]).upper() or "PROCESO",
                n(r["x_m"], "x_m"), n(r["y_m"], "y_m"), n(r["tiempo_min_u"], "tiempo_min_u", 0.0),
                int(n(r["puestos"], "puestos", 1)), t(r["grupo"]).upper()))
        elif hoja == "Conexiones":
            bi = t(r["bidireccional"]).lower()
            if bi not in _VERDADERO | _FALSO:
                raise ValueError(f"«bidireccional» debe ser 1/0 o sí/no (valor: {bi!r}).")
            datos.conexiones.append(Conexion(
                0, t(r["origen"]), t(r["destino"]), n(r["distancia_m"], "distancia_m", 0.0),
                n(r["ancho_m"], "ancho_m"), bi in _VERDADERO or bi == ""))
        elif hoja == "Prendas":
            datos.tipos.append(TipoPrenda(t(r["nombre"]), n(r["peso_unidad_kg"], "peso_unidad_kg"),
                                          int(n(r["unidades_por_lote"], "unidades_por_lote")),
                                          n(r["factor_proceso"], "factor_proceso", 1.0)))
        elif hoja == "Rutas":
            datos.rutas.append((t(r["prenda"]), RutaProceso(0, 0, t(r["ruta"]), n(r["fraccion_pct"], "fraccion_pct", 100.0),
                                                           t(r["secuencia"]))))
        else:
            datos.pedidos.append((t(r["prenda"]), Pedido(
                0, 0, int(n(r["cantidad"], "cantidad")), int(n(r["prioridad"], "prioridad", 3)),
                n(r["llegada_min"], "llegada_min", 0.0), n(r["vencimiento_min"], "vencimiento_min", 0.0),
                t(r["origen_datos"]).upper() or "SINTETICO")))
