# Diseñado y Desarrollado por Jaider Esteban Rodríguez Chiribi y Sebastián González León
# Pruebas de los casos de uso con la base en memoria: simulación, cuellos de botella, requerimientos,
# alertas, cierres de área, inventario y compras, importación y exportación, y proyecto .rlr.

import os
import sqlite3
import zipfile

import pytest

from textilflow.application.errors import ErrorDeNegocio
from textilflow.container import construir
from textilflow.domain.entities import (Conexion, Escenario, Estacion, Pedido, Robot, RutaProceso, Taller, TipoPrenda)
from textilflow.infrastructure.persistence.database import Database


def escenarios(app, taller):
    return app.escenarios.listar(taller.id)


def test_demo_carga_el_caso_experimental(demo):
    app, t = demo
    assert len(app.layout.listar_estaciones(t.id)) == 9
    assert len(app.layout.listar_conexiones(t.id)) == 16
    assert len(app.procesos.listar_tipos()) == 4 and len(app.procesos.listar_rutas(t.id)) == 8
    assert len(app.procesos.listar_pedidos(t.id)) == 5
    s0, s1 = escenarios(app, t)
    assert app.indicadores.execute(s0.id).registros and app.indicadores.execute(s1.id).registros


def test_simulacion_detecta_bordado_como_cuello_de_botella(demo):
    app, t = demo
    s0 = escenarios(app, t)[0]
    r = app.simulacion.execute(s0.id, guardar=False)
    peor = max(r.estados, key=lambda e: e.carga_ofrecida_pct or 0)
    assert peor.codigo == "B" and peor.carga_ofrecida_pct > 85
    assert any(a.codigo == "SIM-CUELLO" and a.elemento == "B" for a in r.alertas)
    assert r.ind.lotes == r.ind.lotes_completados + r.ind.lotes_fallidos + sum(1 for l in r.proceso.lotes if l.estado == "en curso")


def test_simulacion_es_determinista_y_reemplaza_registros(demo):
    app, t = demo
    s0 = escenarios(app, t)[0]
    a = app.simulacion.execute(s0.id)
    b = app.simulacion.execute(s0.id)
    assert len(b.registros) == len(a.registros) == len(app.indicadores.execute(s0.id).registros)
    assert a.indicadores == b.indicadores


def test_prioridad_baja_el_lead_time_del_pedido_urgente(demo):
    app, t = demo
    s0 = escenarios(app, t)[0]
    prio = app.simulacion.execute(s0.id, guardar=False)
    s0.politica = "FIFO"
    app.escenarios.guardar(s0)
    fifo = app.simulacion.execute(s0.id, guardar=False)
    assert prio.ind.lead_time_por_prioridad[1] < fifo.ind.lead_time_por_prioridad[1]


def test_comparacion_usa_calculos_y_kpis_de_proceso(demo):
    app, t = demo
    s0, s1 = escenarios(app, t)
    r = app.comparacion.execute(s0.id, s1.id)
    assert r.comparable and len(r.filas) == 6 and len(r.filas_proceso) == 7
    assert r.series_base["completados"][-1] > 0
    with pytest.raises(ErrorDeNegocio):
        app.comparacion.execute(s0.id, s0.id)


def test_replicas_estadisticas(demo):
    app, t = demo
    s1 = escenarios(app, t)[1]
    r0 = app.replicas.execute(s1.id, 5)
    assert r0.nota and r0.resumen["makespan_min"].desv == pytest.approx(0)
    s1.variabilidad_pct = 20
    app.escenarios.guardar(s1)
    r = app.replicas.execute(s1.id, 12)
    assert not r.nota and r.resumen["makespan_min"].n == 12
    assert r.resumen["makespan_min"].desv > 0 and r.resumen["makespan_min"].ic_inf < r.resumen["makespan_min"].ic_sup
    assert [f["semilla"] for f in r.filas] == list(range(s1.semilla, s1.semilla + 12))
    with pytest.raises(ErrorDeNegocio):
        app.replicas.execute(s1.id, 1)


def test_requerimientos(demo):
    app, t = demo
    s0 = escenarios(app, t)[0]
    r = app.requerimientos.execute(s0.id, 5)
    assert [f.recursos for f in r.barrido] == [1, 2, 3, 4, 5]
    assert r.barrido[0].utilizacion_pct > r.barrido[-1].utilizacion_pct
    assert r.recomendado == 2
    assert any(c.tipo == "Cuello de botella" for c in r.criticos)
    assert any(c.tipo == "Pasillo" and "CO–P" in c.elemento for c in r.criticos)
    camisa = next(v for v in r.viajes_por_prenda if v["prenda"] == "Camisa")
    assert camisa["unidades_por_viaje"] == 100 and camisa["viajes_por_traslado"] == 1
    assert all(x.alcanzable for x in r.ramas)
    assert any(c.estado == "INCUMPLE" for c in r.cumplimientos)
    assert r.recorridos["viajes"] > 0 and r.recorridos["distancia_total_m"] > 0


def test_viajes_aumentan_si_el_equipo_carga_menos(demo):
    app, t = demo
    s0 = escenarios(app, t)[0]
    base = {v["prenda"]: v for v in app.requerimientos.execute(s0.id, 2).viajes_por_prenda}
    s0.capacidad_kg = 5.0
    app.escenarios.guardar(s0)
    chico = {v["prenda"]: v for v in app.requerimientos.execute(s0.id, 2).viajes_por_prenda}
    assert chico["Chaqueta"]["viajes_por_traslado"] > base["Chaqueta"]["viajes_por_traslado"]


def test_alertas_y_parametros_editables(demo):
    app, t = demo
    codigos = {a.codigo for a in app.alertas.execute(t.id)}
    assert {"SIM-CUELLO", "SIM-WIP", "NOR-ancho_min_pasillo_m"} <= codigos
    app.parametros.cambiar_valor("umbral_cuello_botella_pct", 99)
    despues = {(a.codigo, a.elemento) for a in app.alertas.execute(t.id)}
    assert ("SIM-CUELLO", "B") not in despues
    with pytest.raises(ErrorDeNegocio):
        app.parametros.cambiar_valor("no_existe", 1)
    app.parametros.restaurar()
    assert app.parametros.valores()["umbral_cuello_botella_pct"] == 85


def test_rutas_por_nodo_segun_pedido(demo):
    app, t = demo
    r = app.simulacion.execute(escenarios(app, t)[0].id, guardar=False)
    todas = app.rutas.execute(r, "B")
    assert todas and all("B" in x["camino"] for x in todas)
    peds = {x["pedido_id"] for x in todas}
    uno = app.rutas.execute(r, "B", min(peds))
    assert {x["pedido_id"] for x in uno} == {min(peds)}
    assert not app.rutas.execute(r, "B", 999)


def test_layout_no_rompe_relaciones(demo):
    app, t = demo
    est = {e.codigo: e for e in app.layout.listar_estaciones(t.id)}
    with pytest.raises(ErrorDeNegocio):
        app.layout.eliminar_estacion(est["B"].id)
    with pytest.raises(ErrorDeNegocio):
        app.layout.guardar_estacion(Estacion(t.id, "A", "dup", "PROCESO", 0, 0, 1, 1))
    with pytest.raises(ErrorDeNegocio):
        app.layout.guardar_estacion(Estacion(t.id, "A2", "otro almacén", "ALMACEN", 0, 0))
    est["B"].codigo = "BB"
    with pytest.raises(ErrorDeNegocio):
        app.layout.guardar_estacion(est["B"])
    con = app.layout.guardar_conexion(Conexion(t.id, "A", "D", 0, 1.5))
    assert con.distancia_m == pytest.approx(30.02, abs=0.01)


def test_procesos_validan_rutas_y_pedidos(demo):
    app, t = demo
    camisa = next(x for x in app.procesos.listar_tipos() if x.nombre == "Camisa")
    with pytest.raises(ErrorDeNegocio):
        app.procesos.guardar_ruta(RutaProceso(t.id, camisa.id, "mala", 100, "C>CO>D"))
    with pytest.raises(ErrorDeNegocio):
        app.procesos.guardar_ruta(RutaProceso(t.id, camisa.id, "mala", 100, "A>ZZ>D"))
    ok = app.procesos.guardar_ruta(RutaProceso(t.id, camisa.id, "directa", 5, " A > E > D "))
    assert app.procesos.guardar_ruta(RutaProceso(t.id, camisa.id, "grupo", 5, "A>C>@costura>I>E>D")).secuencia == "A>C>@COSTURA>I>E>D"
    assert ok.secuencia == "A>E>D"
    with pytest.raises(ErrorDeNegocio):
        app.procesos.guardar_pedido(Pedido(t.id, 9999, 10))
    with pytest.raises(ErrorDeNegocio):
        app.procesos.eliminar_tipo(camisa.id)
    with pytest.raises(ErrorDeNegocio):
        app.procesos.guardar_tipo(TipoPrenda("camisa", 1, 10))


def test_escenario_amr_toma_parametros_del_robot(demo):
    app, t = demo
    robot = app.robots.listar()[0]
    esc = app.escenarios.guardar(Escenario(t.id, "S1b", "AMR", 3, 480, 9.9, 999, 0, 0, "A", True, robot.id))
    assert (esc.velocidad_ms, esc.capacidad_kg) == (robot.velocidad_ms, robot.capacidad_kg)


def test_compatibilidad_robot_usa_holgura_y_prendas(demo):
    app, t = demo
    robot = app.robots.listar()[0]
    c = app.robots.evaluar_compatibilidad(t.id, robot.id)
    assert c.ancho_requerido_m == pytest.approx(1.30) and not c.compatible_ancho and c.compatible_carga
    assert any("CO–P" in e for e in c.arcos_estrechos)
    pesado = app.robots.guardar(Robot("Débil", 0.5, 0.7, 0.8))
    assert not app.robots.evaluar_compatibilidad(t.id, pesado.id).compatible_carga


def test_validacion_integral_del_demo(demo):
    app, t = demo
    h = {x.codigo: x for x in app.validacion.execute(t.id)}
    assert h["LAY-02"].estado == "PASS" and h["PRO-02"].estado == "PASS" and h["PRO-05"].estado == "PASS"
    assert h["ESC-03"].estado == "PASS" and h["PED-02"].estado == "WARNING" and h["ROB-02"].estado == "WARNING"


def test_validacion_detecta_fallos(app):
    t = app.talleres.guardar(Taller("vacío"))
    h = {x.codigo: x for x in app.validacion.execute(t.id)}
    assert h["LAY-01"].estado == "FAIL" and h["LAY-02"].estado == "FAIL" and h["PED-01"].estado == "FAIL"


def test_simular_sin_datos_da_mensaje_claro(app):
    t = app.talleres.guardar(Taller("t"))
    app.layout.guardar_estacion(Estacion(t.id, "A", "", "ALMACEN", 0, 0))
    app.layout.guardar_estacion(Estacion(t.id, "D", "", "DESPACHO", 5, 0))
    e = app.escenarios.guardar(Escenario(t.id, "S0", "MANUAL", 1, 60, 1, 10, 0, 0, "A"))
    with pytest.raises(ErrorDeNegocio, match="conexiones"):
        app.simulacion.execute(e.id)


def test_eliminar_taller_borra_en_cascada(demo):
    app, t = demo
    app.talleres.eliminar(t.id)
    assert not app.layout.listar_estaciones(t.id) and not app.procesos.listar_rutas(t.id)
    assert not app.procesos.listar_pedidos(t.id) and not app.escenarios.listar(t.id)
    assert app.db.conn.execute("SELECT count(*) FROM registro").fetchone()[0] == 0


def test_excel_ida_y_vuelta(demo, tmp_path):
    app, t = demo
    ruta = str(tmp_path / "taller.xlsx")
    app.datos.exportar_taller(t.id, ruta)
    nuevo, resumen = app.datos.importar_taller(ruta)
    assert "9 nodos" in resumen and "8 ramas" in resumen and "5 pedidos" in resumen
    assert [(e.codigo, e.tiempo_proceso_min, e.capacidad, e.grupo) for e in app.layout.listar_estaciones(nuevo.id)] == \
           [(e.codigo, e.tiempo_proceso_min, e.capacidad, e.grupo) for e in app.layout.listar_estaciones(t.id)]
    assert len(app.procesos.listar_tipos()) == 4
    assert [r.secuencia for r in app.procesos.listar_rutas(nuevo.id)] == [r.secuencia for r in app.procesos.listar_rutas(t.id)]


def test_plantilla_vacia_no_importa(app, tmp_path):
    ruta = str(tmp_path / "p.xlsx")
    app.datos.crear_plantilla(ruta)
    with pytest.raises(ErrorDeNegocio, match="no tiene datos"):
        app.datos.importar_taller(ruta)


def test_importacion_atomica_con_errores(demo, tmp_path):
    from openpyxl import load_workbook
    app, t = demo
    ruta = str(tmp_path / "t.xlsx")
    app.datos.exportar_taller(t.id, ruta)
    wb = load_workbook(ruta)
    wb["Rutas"].append(["Camisa", "rota", 10, "A>NOEXISTE>D"])
    wb["Pedidos"].append(["Fantasma", 10, 1, 0, 0, "SINTETICO"])
    wb["Nodos"].append(["Z", "x", "PROCESO", 1, 1, 0.0, 1, ""])
    wb.save(ruta)
    antes = len(app.talleres.listar())
    with pytest.raises(ErrorDeNegocio) as e:
        app.datos.importar_taller(ruta)
    assert "NOEXISTE" in str(e.value) and "Fantasma" in str(e.value)
    assert len(app.talleres.listar()) == antes


def test_excel_invalido(app, tmp_path):
    ruta = tmp_path / "x.xlsx"
    ruta.write_text("no soy excel")
    with pytest.raises(ErrorDeNegocio):
        app.datos.importar_taller(str(ruta))


def test_proyecto_rlr_ida_y_vuelta(demo, tmp_path):
    app, t = demo
    ruta = str(tmp_path / "p.rlr")
    app.parametros.cambiar_valor("umbral_espera_cola_min", 12)
    app.datos.exportar_proyecto(ruta)
    otra = construir(":memory:")
    otra.datos.restaurar_proyecto(ruta)
    t2 = otra.talleres.listar()[0]
    assert len(otra.procesos.listar_pedidos(t2.id)) == 5 and len(otra.escenarios.listar(t2.id)) == 2
    assert otra.parametros.valores()["umbral_espera_cola_min"] == 12
    a = app.simulacion.execute(escenarios(app, t)[0].id, guardar=False).indicadores
    b = otra.simulacion.execute(otra.escenarios.listar(t2.id)[0].id, guardar=False).indicadores
    assert a == b
    otra.db.close()


def test_restaurar_rechaza_archivos_malos(demo, tmp_path):
    import json
    app, t = demo
    malo = tmp_path / "malo.rlr"
    malo.write_text("no es zip")
    with pytest.raises(ErrorDeNegocio):
        app.datos.restaurar_proyecto(str(malo))
    v1 = tmp_path / "v1.rlr"
    with zipfile.ZipFile(v1, "w") as z:
        z.writestr("proyecto.json", json.dumps({"formato": "TEXTILFLOW-RLR", "version": 1, "tablas": {}}))
    with pytest.raises(ErrorDeNegocio, match="no soportada"):
        app.datos.restaurar_proyecto(str(v1))
    roto = tmp_path / "roto.rlr"
    ruta_ok = str(tmp_path / "ok.rlr")
    app.datos.exportar_proyecto(ruta_ok)
    with zipfile.ZipFile(ruta_ok) as z:
        datos = json.loads(z.read("proyecto.json"))
    datos["tablas"]["pedido"][0]["tipo_prenda_id"] = 12345
    with zipfile.ZipFile(roto, "w") as z:
        z.writestr("proyecto.json", json.dumps(datos))
    with pytest.raises(ErrorDeNegocio, match="No se modificó"):
        app.datos.restaurar_proyecto(str(roto))
    assert len(app.procesos.listar_pedidos(t.id)) == 5


def test_exportar_reporte(demo, tmp_path):
    from openpyxl import load_workbook
    app, t = demo
    s0, s1 = escenarios(app, t)
    x, c = str(tmp_path / "r.xlsx"), str(tmp_path / "r.csv")
    app.reporte.execute(s0.id, s1.id, x, "xlsx")
    app.reporte.execute(s0.id, s1.id, c, "csv")
    ws = load_workbook(x).active
    assert ws.max_row >= 16 and "Lead time" in open(c, encoding="utf-8-sig").read()


def test_base_v1_se_respalda_y_no_se_destruye(tmp_path):
    ruta = str(tmp_path / "vieja.db")
    con = sqlite3.connect(ruta)
    con.executescript("CREATE TABLE solicitud(id INTEGER); INSERT INTO solicitud VALUES (1);")
    con.commit(); con.close()
    db = Database(ruta)
    assert db.respaldo_v1 and os.path.exists(db.respaldo_v1)
    assert sqlite3.connect(db.respaldo_v1).execute("SELECT count(*) FROM solicitud").fetchone()[0] == 1
    assert db.conn.execute("PRAGMA user_version").fetchone()[0] == 2
    db.close()
    assert Database(ruta).respaldo_v1 == ""


def test_ejemplo_xlsx_se_importa_y_simula():
    from pathlib import Path
    app = construir(":memory:")
    ruta = Path(__file__).resolve().parents[1] / "ejemplos" / "01_confeccion_andina.xlsx"
    t, _ = app.datos.importar_taller(str(ruta))
    assert app.demo.preparar_escenario_base(t) is True
    assert app.demo.preparar_escenario_base(t) is False
    s0 = app.escenarios.listar(t.id)[0]
    assert app.indicadores.execute(s0.id).registros


def _cerrar(app, taller, codigo, **kw):
    from textilflow.domain.entities import Cierre
    return app.cierres.guardar(Cierre(taller.id, codigo, **kw))


def _esc0(app, taller):
    return app.escenarios.listar(taller.id)[0]


def test_cierre_desvia_el_trabajo_al_area_gemela(demo):
    app, taller = demo
    esc = _esc0(app, taller)
    antes = app.simulacion.execute(esc.id, False)
    _cerrar(app, taller, "CO", desde_min=0, hasta_min=0)
    despues = app.simulacion.execute(esc.id, False)
    visitas = lambda r, a: sum(1 for v in r.proceso.visitas if v.estacion == a and v.inicio_min is not None)
    assert visitas(antes, "CO") > 0
    assert visitas(despues, "CO") == 0 and visitas(despues, "CO2") > visitas(antes, "CO2")
    assert despues.ind.lotes_fallidos == 0


def test_cierre_de_area_unica_falla_los_lotes_con_motivo(demo):
    app, taller = demo
    esc = _esc0(app, taller)
    _cerrar(app, taller, "B", motivo="MANTENIMIENTO")
    r = app.simulacion.execute(esc.id, False)
    fallidos = [l for l in r.proceso.lotes if l.estado == "fallido"]
    assert fallidos and all("Bordado" in l.motivo or "«B»" in l.motivo for l in fallidos)


def test_reapertura_reanuda_el_proceso(demo):
    app, taller = demo
    esc = _esc0(app, taller)
    _cerrar(app, taller, "B", desde_min=0, hasta_min=120)
    r = app.simulacion.execute(esc.id, False)
    inicios = [v.inicio_min for v in r.proceso.visitas if v.estacion == "B" and v.inicio_min is not None]
    assert inicios and min(inicios) >= 120 - 1e-6


def test_no_se_puede_cerrar_almacen_ni_despacho(demo):
    from textilflow.application.errors import ErrorDeNegocio
    app, taller = demo
    for cod in ("A", "D"):
        with pytest.raises(ErrorDeNegocio):
            _cerrar(app, taller, cod)


def test_informe_de_cierre(demo):
    from textilflow.application.errors import ErrorDeNegocio
    app, taller = demo
    esc = _esc0(app, taller)
    with pytest.raises(ErrorDeNegocio):
        app.informe_cierres.execute(esc.id)
    _cerrar(app, taller, "CO", motivo="FALLA_MAQUINARIA", detalle="Se quemó el motor", desde_min=60)
    inf = app.informe_cierres.execute(esc.id)
    assert inf.nivel in ("ALTO", "MEDIO", "BAJO") and inf.rutas and inf.recomendaciones
    assert any(r[0] == "COSTURA" for r in inf.reasignaciones)


def test_cierres_viajan_en_el_proyecto(demo, tmp_path):
    app, taller = demo
    _cerrar(app, taller, "CO", detalle="x")
    ruta = str(tmp_path / "p.rlr")
    app.datos.exportar_proyecto(ruta)
    otra = construir(":memory:")
    otra.datos.restaurar_proyecto(ruta)
    t2 = otra.talleres.listar()[0]
    assert otra.cierres.cerrados(t2.id) == frozenset({"CO"})
    otra.db.close()


def test_requerimientos_respetan_el_cierre(demo):
    app, taller = demo
    esc = _esc0(app, taller)
    antes = app.requerimientos.execute(esc.id)
    _cerrar(app, taller, "CO")
    despues = app.requerimientos.execute(esc.id)
    assert antes is not None and despues is not None


def test_exportar_contingencias_a_excel(demo, tmp_path):
    from openpyxl import load_workbook
    from textilflow.application.errors import ErrorDeNegocio
    app, taller = demo
    esc = _esc0(app, taller)
    ruta = str(tmp_path / "c.xlsx")
    with pytest.raises(ErrorDeNegocio):
        app.exportar_contingencias.execute(taller.id, esc.id, ruta)
    _cerrar(app, taller, "CO", detalle="Se quemó el motor", desde_min=60)
    assert app.exportar_contingencias.execute(taller.id, esc.id, ruta) == 1
    wb = load_workbook(ruta)
    hoja = wb["Contingencias"]
    cab = [c.value for c in hoja[3]]
    fila = {k: v for k, v in zip(cab, [c.value for c in hoja[4]])}
    assert fila["Área"] == "Costura 1" and fila["Motivo"] == "Falla de maquinaria" and fila["Fecha y hora"][:2] == "20"
    assert wb["Rutas alternativas"]["D4"].value.count("CO2") == 1


def test_cache_de_simulacion_se_reutiliza_y_se_invalida_al_cambiar_los_datos(demo):
    from textilflow.domain.entities import Cierre
    app, taller = demo
    esc = app.escenarios.listar(taller.id)[0]
    a = app.simulacion.execute(esc.id, False)
    b = app.simulacion.execute(esc.id, False)
    assert a is b
    app.cierres.guardar(Cierre(taller.id, "C", "FALLA_MAQUINARIA", "", 0.0, 0.0, True, True))
    c = app.simulacion.execute(esc.id, False)
    assert c is not a
    assert sum(l.estado == "fallido" for l in c.proceso.lotes) > sum(l.estado == "fallido" for l in a.proceso.lotes)
    d = app.simulacion.execute(esc.id, True)
    assert d is not c and c.registros is not d.registros


def test_toda_alerta_trae_posible_solucion(demo):
    from textilflow.domain.alertas import Alerta, TIPO_CUELLO, cuellos_de_botella
    app, taller = demo
    _cerrar(app, taller, "CO", desde_min=0, hasta_min=0)
    alertas = app.alertas.execute(taller.id)
    assert alertas
    sin = [a.codigo for a in alertas if not a.solucion]
    assert not sin, sin
    assert {"CIERRE", "SIM-CUELLO"} <= {a.codigo for a in alertas}
    cuellos = cuellos_de_botella(alertas)
    assert cuellos and all(a.codigo in TIPO_CUELLO for a in cuellos)
    niveles = [a.nivel for a in cuellos]
    orden = ["CRÍTICA", "ALTA", "MEDIA", "INFO"]
    assert niveles == sorted(niveles, key=orden.index)


def test_alertas_de_configuracion_traen_solucion(app):
    from textilflow.domain.alertas import alertas_estaticas
    alertas = alertas_estaticas([], [], [], [], [])
    assert {a.codigo for a in alertas} == {"CFG-ALM", "CFG-DES"}
    assert all(a.solucion for a in alertas)


def _tipo(app, nombre):
    return next(t for t in app.procesos.listar_tipos() if t.nombre == nombre)


def test_demo_trae_stock_y_alerta_de_camisas_que_no_alcanzan(demo):
    app, taller = demo
    camisa = _tipo(app, "Camisa")
    s = next(x for x in app.inventario.listar_stock(taller.id) if x.tipo_prenda_id == camisa.id)
    b = app.inventario.balance(taller.id, s)
    assert (s.unidades, b["demanda"], b["falta"]) == (400, 550, 150)
    assert b["sugerida"] == 200
    a = next(x for x in app.alertas.execute(taller.id) if x.codigo == "STOCK-INSUF")
    assert a.elemento == "Camisa" and "200" in a.solucion and a.solucion


def test_taller_sin_unidades_genera_alerta_critica_y_aviso_de_compra(demo):
    from textilflow.domain.entities import Stock
    app, taller = demo
    pant = _tipo(app, "Pantalón")
    s = next(x for x in app.inventario.listar_stock(taller.id) if x.tipo_prenda_id == pant.id)
    app.inventario.guardar_stock(Stock(taller.id, pant.id, 0, 30, s.id))
    a = next(x for x in app.alertas.execute(taller.id) if x.codigo == "STOCK-AGOTADO")
    assert a.nivel == "CRÍTICA" and a.elemento == "Pantalón"
    aviso = app.inventario.generar_aviso(taller.id, pant.id)
    assert aviso.estado == "PENDIENTE" and aviso.cantidad == 230
    codigos = {x.codigo for x in app.alertas.execute(taller.id)}
    assert "COMPRA-PENDIENTE" in codigos
    a = next(x for x in app.alertas.execute(taller.id) if x.codigo == "STOCK-AGOTADO")
    assert a.nivel == "MEDIA" and "en compra" in a.detalle
    with pytest.raises(ErrorDeNegocio):
        app.inventario.generar_aviso(taller.id, pant.id)
    texto = app.inventario.aviso_para_compras(taller.id)
    assert "Pantalón: 230 unidades" in texto and taller.nombre in texto


def test_recibir_la_compra_suma_al_stock_y_no_se_puede_modificar(demo):
    app, taller = demo
    camisa = _tipo(app, "Camisa")
    aviso = app.inventario.generar_aviso(taller.id, camisa.id)
    aviso.estado = "PEDIDA"
    app.inventario.guardar_compra(aviso)
    aviso.estado = "RECIBIDA"
    app.inventario.guardar_compra(aviso)
    s = next(x for x in app.inventario.listar_stock(taller.id) if x.tipo_prenda_id == camisa.id)
    assert s.unidades == 600
    assert not [a for a in app.alertas.execute(taller.id) if a.codigo.startswith(("STOCK", "COMPRA"))]
    aviso.estado = "PENDIENTE"
    with pytest.raises(ErrorDeNegocio):
        app.inventario.guardar_compra(aviso)


def test_stock_una_fila_por_prenda_y_validaciones(demo):
    from textilflow.domain.entities import Stock
    app, taller = demo
    camisa = _tipo(app, "Camisa")
    n = len(app.inventario.listar_stock(taller.id))
    app.inventario.guardar_stock(Stock(taller.id, camisa.id, 10, 5))
    assert len(app.inventario.listar_stock(taller.id)) == n
    with pytest.raises(ErrorDeNegocio):
        app.inventario.guardar_stock(Stock(taller.id, camisa.id, -1, 0))
    with pytest.raises(ErrorDeNegocio):
        app.inventario.generar_aviso(taller.id, 99999)


def test_inventario_viaja_en_el_proyecto(demo, tmp_path):
    app, taller = demo
    ruta = str(tmp_path / "p.rlr")
    app.datos.exportar_proyecto(ruta)
    app.talleres.eliminar(taller.id)
    assert not app.inventario.listar_stock(taller.id)
    app.datos.restaurar_proyecto(ruta)
    assert len(app.inventario.listar_stock(taller.id)) == 4
